#!/usr/bin/env python3
"""Every-sale-matters scoreboard pipeline.

Downloads the team's Google Sheets (public "anyone with link" access),
applies the same counting rules as the comparison report, and writes
scoreboard.json for the widget.

Counting rules (from the scoreboard / comparison report Method tab):
- A registration = a named registrar row with a student name OR CRM ID.
- Status "Cancelled" excluded; "Confirmed" and "Docs Received" count.
- August baseline: person's August total / 20 recorded days.
- Expected for a period = August per-day * recorded days in that period.
- Personal best = best single day since (and including) August.
- Christa and ChristaB are separate people; ClarenceB excluded.
"""

import json
import re
import sys
import urllib.request
from datetime import datetime, date
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
CACHE = BASE / "sheets"
OUT = BASE / "scoreboard.json"

SHEETS = {
    "current": "1KBcTzi88HrO1kVoTHOvM8ZRv52oOjgE-kPIgIg22J_A",   # current month's source sheet
    "previous": "1HEe6HTjmz1knhvp02DliHapCqtM0pZuJasK9pYXb034",  # previous month (August)
    "incentive": "1A8aqIF5etGw2CelyXDqCfnSlgCshGw3a0VqV80CvfL8",  # weekly incentives + cash
    "comparison": "1j0DhPZkcBRQgGxKtbEOWNSKNXkqqhUh16BWgYeHTst8", # roster + colleges + Method
}

AUG_RECORDED_DAYS = 20  # per the Method tab

ENROLMENT_LEVELS = [  # (min qualifying debit orders, rand amount)
    (25, 400), (30, 550), (35, 650), (40, 750), (45, 850), (50, 1000),
    (55, 1100), (60, 1200), (65, 1300), (70, 1400), (75, 1500),
    (80, 1600), (85, 1700), (90, 1800), (95, 1900), (100, 2000),
]

CASH_BRACKETS = [  # (min amount, incentive) — applied per row/payment
    (40000, 1000), (30000, 700), (20000, 500), (10000, 300),
    (5000, 200), (2500, 100), (700, 50), (0, 0),
]

DAY_TAB_RE = re.compile(r"^\s*(\d{1,2})\s*([A-Za-z]{3,})\s*$")


def month_num(mon):
    months = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
              "jul": 7, "aug": 8, "sept": 9, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
    return months.get(mon.lower())


def parse_day_tab(title):
    m = DAY_TAB_RE.match(title.strip())
    if not m:
        return None
    n = month_num(m.group(2))
    if not n:
        return None
    return date(2026, n, int(m.group(1)))


def download(sheet_id, dest):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as r:
        data = r.read()
    dest.write_bytes(data)
    return dest


def norm(name):
    return re.sub(r"\s+", " ", str(name or "").strip()).lower()


def g(row, i):
    """Safe column access: '' for missing/short rows."""
    return row[i] if i < len(row) and row[i] is not None else ""


def count_daily_tab(ws):
    """Return {registrar_norm: count} of qualifying registrations in one daily tab."""
    counts = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        registrar = str(g(r, 0)).strip()
        if not registrar or registrar.startswith("="):
            continue
        crm = str(g(r, 2)).strip()
        student = str(g(r, 4)).strip()
        status = str(g(r, 12)).strip().lower()
        if not (crm or student):
            continue  # no student name and no CRM ID
        if status == "cancelled":
            continue
        key = norm(registrar)
        counts[key] = counts.get(key, 0) + 1
    return counts


def count_daily_debit_orders(ws):
    """{registrar_norm: count} of qualifying D/O ('X' in col L) rows."""
    counts = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        registrar = str(g(r, 0)).strip()
        if not registrar or registrar.startswith("="):
            continue
        crm = str(g(r, 2)).strip()
        student = str(g(r, 4)).strip()
        status = str(g(r, 12)).strip().lower()
        do = str(g(r, 11)).strip().upper()
        if not (crm or student) or status == "cancelled":
            continue
        if do != "X":
            continue
        key = norm(registrar)
        counts[key] = counts.get(key, 0) + 1
    return counts


def count_incentive_do(ws):
    """Incentive-sheet daily tabs have a different layout: A=Registrar, E=D/O, F=Comments."""
    counts = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        registrar = str(g(r, 0)).strip()
        if not registrar or registrar.startswith("="):
            continue
        do = str(g(r, 4)).strip().upper()
        status = str(g(r, 5)).strip().lower()
        if do != "X" or status == "cancelled":
            continue
        key = norm(registrar)
        counts[key] = counts.get(key, 0) + 1
    return counts


def load_daily_book(path, do_counter=count_daily_debit_orders):
    """{date: {'reg': {name: n}, 'do': {name: n}}} for every day tab in a workbook."""
    wb = openpyxl.load_workbook(path, read_only=True)
    days = {}
    for ws in wb.worksheets:
        d = parse_day_tab(ws.title)
        if not d:
            continue
        days[d] = {"reg": count_daily_tab(ws), "do": do_counter(ws)}
    wb.close()
    return dict(sorted(days.items()))


def cash_incentive_for(amount):
    for floor, inc in CASH_BRACKETS:
        if amount >= floor:
            return inc
    return 0


def enrolment_incentive_for(n):
    best = 0
    for floor, amt in ENROLMENT_LEVELS:
        if n >= floor:
            best = amt
    return best


def next_enrolment_level(n):
    for floor, amt in ENROLMENT_LEVELS:
        if n < floor:
            return floor, amt
    return None, None


def week_of(d):
    """Monday of the work week containing d."""
    return d.fromordinal(d.toordinal() - d.weekday())


def main():
    CACHE.mkdir(exist_ok=True)
    files = {}
    for key, sid in SHEETS.items():
        dest = CACHE / f"{key}_{sid}.xlsx"
        try:
            download(sid, dest)
            files[key] = dest
            print(f"downloaded {key}: {dest.stat().st_size} bytes", file=sys.stderr)
        except Exception as e:
            if dest.exists():
                files[key] = dest
                print(f"download failed for {key}, using cache: {e}", file=sys.stderr)
            else:
                raise

    # ---- roster from comparison report ----
    wb = openpyxl.load_workbook(files["comparison"], read_only=True)
    ws = wb["All Colleges"]
    roster = []  # {name, college, aug_total, aug_per_day}
    for r in ws.iter_rows(min_row=9, values_only=True):
        name = str(g(r, 1)).strip()
        if not name or name.upper() == "TOTAL" or name.lower() == "registrar":
            continue
        roster.append({
            "name": name,
            "college": str(g(r, 2)).strip(),
            "aug_total": float(g(r, 3) or 0),
            "aug_per_day": float(g(r, 4) or 0),
        })
    wb.close()
    roster_keys = {norm(p["name"]) for p in roster}

    # ---- daily data ----
    prev_days = load_daily_book(files["previous"])   # August
    curr_days = load_daily_book(files["current"])    # current month
    inc_days = load_daily_book(files["incentive"], do_counter=count_incentive_do)  # incentive sheet's own daily tabs

    all_days = {**prev_days, **curr_days}
    if not curr_days:
        print("WARNING: no daily tabs in current month sheet", file=sys.stderr)
        curr_days = prev_days
    latest = max(curr_days)
    today_counts = curr_days[latest]["reg"]

    day_list = sorted(curr_days)
    yesterday = day_list[-2] if len(day_list) >= 2 else None

    def period_days(period):
        if period == "today":
            return [latest]
        if period == "yesterday":
            return [yesterday] if yesterday else []
        if period == "week":
            wk = week_of(latest)
            return [d for d in day_list if week_of(d) == wk]
        if period == "month":
            return day_list
        return []

    def sum_period(days, source):
        total = {}
        for d in days:
            for k, v in source.get(d, {"reg": {}})["reg"].items():
                total[k] = total.get(k, 0) + v
        return total

    periods = {}
    for p in ("today", "yesterday", "week", "month"):
        days = period_days(p)
        counts = sum_period(days, curr_days)
        periods[p] = {
            "recorded_days": len(days),
            "days": [d.isoformat() for d in days],
            "counts": {k: counts.get(k, 0) for k in roster_keys},
        }

    # ---- personal best since August (single day) ----
    best = {k: 0 for k in roster_keys}
    prev_best = {k: 0 for k in roster_keys}
    for d, data in all_days.items():
        for k, v in data["reg"].items():
            if d <= latest and k in best:
                best[k] = max(best[k], v)
            if d < latest and k in prev_best:
                prev_best[k] = max(prev_best[k], v)

    # ---- colleges: current month vs August pace ----
    month_days_n = len(periods["month"]["days"])
    colleges = {}
    for p in roster:
        c = p["college"]
        e = colleges.setdefault(c, {"now": 0, "expected": 0.0, "aug_per_day": 0.0})
        e["now"] += periods["month"]["counts"].get(norm(p["name"]), 0)
        e["expected"] += p["aug_per_day"] * month_days_n
        e["aug_per_day"] += p["aug_per_day"]
    for c, e in colleges.items():
        e["pct"] = (e["now"] / e["expected"] * 100) if e["expected"] else 0.0

    # ---- weekly incentives ----
    wk = week_of(latest)
    inc_week_days = [d for d in sorted(inc_days) if week_of(d) == wk]
    do_counts = {}
    for d in inc_week_days:
        for k, v in inc_days[d]["do"].items():
            do_counts[k] = do_counts.get(k, 0) + v

    wb = openpyxl.load_workbook(files["incentive"], read_only=True)
    # find the week-number of the current week from tab names like 'Week 3 Total'
    week_tabs = [t for t in wb.sheetnames if re.match(r"^Week\s+\d+\s+Total$", t.strip())]
    cash_rows = []
    paid = {}
    def week_range_dates(tabname):
        # 'Week 3 Total' -> dates from inc_days grouped by week number order
        return None
    # map week tab to its Monday via the daily tabs it references: simpler —
    # weeks are in order; compute week index of current week among all weeks present
    all_inc_days_sorted = sorted(inc_days)
    week_mondays = sorted({week_of(d) for d in all_inc_days_sorted})
    cur_week_idx = week_mondays.index(wk) if wk in week_mondays else None
    cash_by_person = {}
    ref_by_person = {}
    if cur_week_idx is not None:
        tab = f"Week {cur_week_idx + 1} Total"
        if tab in wb.sheetnames:
            ws = wb[tab]
            for r in ws.iter_rows(min_row=4, values_only=True):
                name = str(g(r, 0)).strip()
                if not name:
                    continue
                paid_flag = str(g(r, 18)).strip().lower()
                paid[norm(name)] = paid_flag in ("true", "yes", "1", "x")
        cash_tab = f"Cash Payments Week {cur_week_idx + 1}"
        if cash_tab in wb.sheetnames:
            ws = wb[cash_tab]
            for r in ws.iter_rows(min_row=3, values_only=True):
                name = str(g(r, 1)).strip()
                amt = g(r, 2)
                if not name or not isinstance(amt, (int, float)):
                    continue
                if isinstance(amt, bool):
                    continue
                amt = float(amt)
                k = norm(name)
                cash_by_person[k] = cash_by_person.get(k, 0) + amt
                ref_by_person[k] = ref_by_person.get(k, 0) + cash_incentive_for(amt)
    wb.close()

    incentives = {}
    for p in roster:
        k = norm(p["name"])
        n_do = do_counts.get(k, 0)
        cash = cash_by_person.get(k, 0.0)
        inc = {
            "debit_orders": n_do,
            "enrolment_incentive": enrolment_incentive_for(n_do),
            "cash_collected": cash,
            "cash_incentive": ref_by_person.get(k, 0),
            "paid": paid.get(k, False),
        }
        inc["total"] = inc["enrolment_incentive"] + inc["cash_incentive"]
        floor, amt = next_enrolment_level(n_do)
        inc["next_level_at"] = floor
        inc["next_level_amount"] = amt
        inc["to_next_level"] = (floor - n_do) if floor else 0
        incentives[k] = inc

    # ---- people payload ----
    people = []
    for p in roster:
        k = norm(p["name"])
        daily_best = best.get(k, 0)
        people.append({
            "name": p["name"],
            "college": p["college"],
            "aug_per_day": p["aug_per_day"],
            "today": periods["today"]["counts"].get(k, 0),
            "yesterday": periods["yesterday"]["counts"].get(k, 0) if yesterday else None,
            "week": sum(periods["week"]["counts"].get(k, 0) for _ in [0]),
            "month": periods["month"]["counts"].get(k, 0),
            "best_day": daily_best,
            "prev_best": prev_best.get(k, 0),
            "incentive": incentives[k],
        })
    for ppl in people:
        ppl["week"] = sum(
            curr_days[d]["reg"].get(norm(ppl["name"]), 0) for d in period_days("week")
        )

    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "latest_day": latest.isoformat(),
        "month_recorded_days": month_days_n,
        "aug_recorded_days": AUG_RECORDED_DAYS,
        "people": people,
        "colleges": colleges,
        "periods": {p: {"recorded_days": v["recorded_days"], "days": v["days"]}
                    for p, v in periods.items()},
    }
    OUT.write_text(json.dumps(payload, indent=1, default=str))
    print(f"wrote {OUT}", file=sys.stderr)
    return payload


if __name__ == "__main__":
    main()
